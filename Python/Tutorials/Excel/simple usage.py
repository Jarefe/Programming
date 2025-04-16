from openpyxl import Workbook
import os

wb = Workbook()
ws = wb.active
treeData = [["Type", "Leaf Color", "Height"], ["Maple", "Red", 549], ["Oak", "Green", 783], ["Pine", "Green", 1204]]

# can use append because data is list of lists
for row in treeData:
    ws.append(row)

# edit font
from openpyxl.styles import Font
ft = Font(bold=True)
for row in ws["A1:C1"]:
    for cell in row:
        cell.font = ft

# create charts
from openpyxl.chart import BarChart, series , Reference

chart = BarChart()
chart.type = "col"
chart.title = "Tree Height"
chart.y_axis.title = 'Height (cm)'
chart.x_axis.title = 'Tree Type'
chart.legend = None

# add references to where data is and pass to chart object
data = Reference(ws, min_col=3, min_row=2, max_row=4, max_col=3)
categories = Reference(ws, min_col=1, min_row=2, max_row=4, max_col=1)
chart.add_data(data)
chart.set_categories(categories)

# add to sheet
ws.add_chart(chart, 'E1')

# get path of script
script_dir = os.path.dirname(os.path.abspath(__file__))

# define path where excel file should be saved
excel_path = os.path.join(script_dir, "Treedata.xlsx")

wb.save(excel_path)
print(f'Workbook saved to: {excel_path}')